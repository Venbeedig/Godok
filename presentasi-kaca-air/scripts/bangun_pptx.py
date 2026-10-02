#!/usr/bin/env python3
"""bangun_pptx.py — susun .pptx (transisi Morph) dari bahan ekspor.cjs.
Latar = gambar diam; kartu kaca = bentuk PowerPoint asli (gradasi transparan);
teks = kotak teks asli yang bisa diedit; mawar/kartu palet/ikon = gambar PNG transparan.
Objek yang sama di dua slide berurutan diberi nama '!!nama' supaya Morph menggerakkannya.
Pakai: python3 scripts/bangun_pptx.py <folder-ekspor> <keluaran.pptx>"""
import json
import pathlib
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

FONT = 'Segoe UI'          # tersedia di semua Windows; ganti 'Poppins' bila font itu terpasang
NSA = 'http://schemas.openxmlformats.org/drawingml/2006/main'
MORPH_MS = 1400

src = pathlib.Path(sys.argv[1])
out = pathlib.Path(sys.argv[2])
tata = json.loads((src / 'tata.json').read_text(encoding='utf-8'))

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
K = 12192000 / tata['lebar']                      # EMU per piksel kanvas


def px(v):
    return Emu(int(round(v * K)))


def pt(v):                                         # 1 px kanvas = 0,5 pt
    return Pt(round(v * 0.5 * 2) / 2)


def campur(hex_a, hex_b, t):                       # t=0 → a, t=1 → b
    a = [int(hex_a.lstrip('#')[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_b.lstrip('#')[i:i + 2], 16) for i in (0, 2, 4)]
    return ''.join(f'{round(x + (y - x) * t):02X}' for x, y in zip(a, b))


def beri_nama(shape, nama, akhiran=''):
    if nama:
        shape.name = '!!' + nama + akhiran


def kosongkan_spPr(shape):
    spPr = shape._element.spPr
    for tag in ('a:solidFill', 'a:gradFill', 'a:noFill', 'a:ln', 'a:effectLst'):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    return spPr


def sisip_setelah_geom(spPr, xmls):
    geom = spPr.find(qn('a:prstGeom'))
    i = list(spPr).index(geom)
    for k, x in enumerate(xmls):
        spPr.insert(i + 1 + k, etree.fromstring(x))


def gradasi(stops, sudut=65):
    gs = ''.join(f'<a:gs pos="{p * 1000}"><a:srgbClr val="{c}"><a:alpha val="{a * 1000}"/></a:srgbClr></a:gs>' for p, c, a in stops)
    return f'<a:gradFill xmlns:a="{NSA}" rotWithShape="1"><a:gsLst>{gs}</a:gsLst><a:lin ang="{sudut * 60000}" scaled="0"/></a:gradFill>'


def garis(alpha=85, lebar_pt=1.75):
    return (f'<a:ln xmlns:a="{NSA}" w="{int(lebar_pt * 12700)}"><a:solidFill><a:srgbClr val="FFFFFF">'
            f'<a:alpha val="{alpha * 1000}"/></a:srgbClr></a:solidFill></a:ln>')


def bayangan(rgb, alpha=26):
    return (f'<a:effectLst xmlns:a="{NSA}"><a:outerShdw blurRad="508000" dist="190500" dir="5400000" algn="t" rotWithShape="0">'
            f'<a:srgbClr val="{rgb}"><a:alpha val="{alpha * 1000}"/></a:srgbClr></a:outerShdw></a:effectLst>')


def bentuk_kaca(slide, it, palet):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(it['x']), px(it['y']), px(it['w']), px(it['h']))
    s.adjustments[0] = min(0.5, it.get('radius', 44) / max(1, min(it['w'], it['h'])))
    if it.get('varian') == 'susu':
        stops = [(0, 'FFFFFF', 84), (55, 'FFFFFF', 60), (100, 'FFFFFF', 70)]
    elif it.get('varian') == 'warna':
        tint = campur(palet['c1'], '#FFFFFF', 0.62)
        stops = [(0, tint, 78), (100, 'FFFFFF', 45)]
    else:
        stops = [(0, 'FFFFFF', 70), (52, 'FFFFFF', 40), (100, 'FFFFFF', 52)]
    bay = ''.join(f'{int(v):02X}' for v in palet['bayang'].split(','))
    sisip_setelah_geom(kosongkan_spPr(s), [gradasi(stops), garis(), bayangan(bay)])
    s.text_frame.text = ''
    beri_nama(s, it['nama'])
    return s


def bentuk_pil(slide, it):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(it['x']), px(it['y']), px(it['w']), px(it['h']))
    s.adjustments[0] = 0.5
    sisip_setelah_geom(kosongkan_spPr(s), [gradasi([(0, 'FFFFFF', 70), (100, 'FFFFFF', 58)], 90), garis(92, 1.25)])
    beri_nama(s, it['nama'], '-pil')


def isi_run(p, r):
    if r['t'] == '\n':
        p.add_line_break()
        return
    run = p.add_run()
    run.text = r['t']
    f = run.font
    f.name = FONT
    f.size = pt(r['ukuran'])
    f.bold = r['tebal'] >= 600
    f.italic = bool(r['miring'])
    f.color.rgb = RGBColor.from_string(r['warna'])
    if abs(r.get('spasi', 0)) > 0.05:
        run._r.get_or_add_rPr().set('spc', str(int(round(r['spasi'] * 50))))


RATA = {'kiri': PP_ALIGN.LEFT, 'tengah': PP_ALIGN.CENTER, 'kanan': PP_ALIGN.RIGHT}


def kotak_teks(slide, it, palet):
    if it.get('pil'):
        bentuk_pil(slide, it)
    tb = slide.shapes.add_textbox(px(it['x']), px(it['y']), px(it['w']), px(it['h']))
    tf = tb.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    pad = it.get('pad') or [0, 0, 0, 0]
    tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = (px(v) for v in pad)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if it.get('tengahV') else MSO_ANCHOR.TOP
    if it['jenis'] == 'daftar':
        tf.word_wrap = True
        for i, b in enumerate(it['butir']):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = RATA[it['rata']]
            p.line_spacing = pt(b['tinggiBaris'])
            p.space_after = pt(b['jarak'])
            pPr = p._p.get_or_add_pPr()
            indent = b['indent'] or -b['gantung']
            if indent:
                pPr.set('marL', str(int(px(indent))))
                pPr.set('indent', str(-int(px(indent))))
            if it.get('peluru'):
                pPr.append(etree.fromstring(f'<a:buClr xmlns:a="{NSA}"><a:srgbClr val="{palet["c1"].lstrip("#").upper()}"/></a:buClr>'))
                pPr.append(etree.fromstring(f'<a:buSzPct xmlns:a="{NSA}" val="90000"/>'))
                pPr.append(etree.fromstring(f'<a:buFont xmlns:a="{NSA}" typeface="Arial"/>'))
                pPr.append(etree.fromstring(f'<a:buChar xmlns:a="{NSA}" char="●"/>'))
            for r in b['par']:
                isi_run(p, r)
    else:
        baris = max(1, round((it['h'] - pad[0] - pad[2]) / max(1, it['tinggiBaris'])))
        tf.word_wrap = not it.get('pil') and (baris > 1 or it['w'] > 900)
        for i, runs in enumerate(it['par']):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = RATA[it['rata']]
            if not it.get('pil'):
                p.line_spacing = pt(it['tinggiBaris'])
            for r in runs:
                isi_run(p, r)
    beri_nama(tb, it['nama'])


def transisi(slide, morph=True):
    if morph:
        xml = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
               '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
               '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               f'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="{MORPH_MS}">'
               '<p159:morph option="byObject"/></p:transition></mc:Choice><mc:Fallback>'
               '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/></p:transition>'
               '</mc:Fallback></mc:AlternateContent>')
    else:
        xml = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/></p:transition>'
    sld = slide._element
    jangkar = sld.find(qn('p:clrMapOvr'))
    if jangkar is None:
        jangkar = sld.find(qn('p:cSld'))
    jangkar.addnext(etree.fromstring(xml))


def perkecil(f, lebar_tampil, kali=1.5):
    """Batasi resolusi gambar ke 1,5x ukuran tampil supaya .pptx tetap ringan."""
    im = Image.open(f)
    target = int(lebar_tampil * kali)
    if im.width <= target:
        return f
    kecil = src / 'kecil' / f.name
    kecil.parent.mkdir(exist_ok=True)
    im.resize((target, round(im.height * target / im.width)), Image.LANCZOS).save(kecil, optimize=True)
    return kecil


kosong = prs.slide_layouts[6]
for n, s in enumerate(tata['slides']):
    slide = prs.slides.add_slide(kosong)
    latar = slide.shapes.add_picture(str(src / s['latar']), 0, 0, prs.slide_width, prs.slide_height)
    latar.name = f'Latar {n + 1}'
    for it in s['items']:
        if it['jenis'] == 'kaca':
            bentuk_kaca(slide, it, s['palet'])
        elif it['jenis'] == 'gambar':
            g = slide.shapes.add_picture(str(perkecil(src / it['file'], it['w'])), px(it['x']), px(it['y']), px(it['w']), px(it['h']))
            beri_nama(g, it['nama'])
        else:
            kotak_teks(slide, it, s['palet'])
    transisi(slide, morph=n > 0)
    slide.notes_slide.notes_text_frame.text = f"{s['bagian']} — ⟨tulis catatan pembicara di sini⟩"

out.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out))
print(f'{out}  ({out.stat().st_size / 1024 / 1024:.1f} MB, {len(tata["slides"])} slide)')
