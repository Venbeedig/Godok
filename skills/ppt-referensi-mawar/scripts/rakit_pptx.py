#!/usr/bin/env python3
"""Rakit .pptx dari PETA SLIDE (JSON), dalam dua mode.

MODE REFERENSI (utama — hasil meniru file referensi LO):
  python3 rakit_pptx.py --peta peta.json --referensi referensi.pptx --keluar hasil.pptx \
                        [--mawar mawar.png] [--pojok]
  Tiap slide di peta memakai "pola": nomor slide referensi yang disalin UTUH (latar, gambar, bentuk,
  font, warna, animasi), lalu teks bentuk bernama diganti lewat "isi": {"<nama bentuk>": "teks" | ["baris", …]}.
  "buang": [nama bentuk] menghapus bentuk; "mawar": {"x":…, "y":…, "w":…} (inci) menempel mawar+logo;
  "ganti_gambar": "<nama bentuk gambar>" menaruh mawar tepat di posisi gambar itu lalu membuang gambarnya.
  --pojok menempel mawar kecil di slide isi; posisinya bisa diatur lewat "pojok": {"x","y","w"} di akar peta.
  Nama bentuk dilihat dari hasil bedah_referensi.py (peta-bentuk.txt). Slide asli referensi dibuang di akhir.

MODE MANDIRI (cadangan bila referensi belum ada):
  python3 rakit_pptx.py --peta peta.json --keluar hasil.pptx [--mawar mawar.png] [--latar-folder latar/] [--aset <folder kc-*>]
  Slide disusun dari "tipe" (sampul, kelompok, anggota, agenda, poin, kartu, banding, kisi, pernyataan, penutup) di atas
  gambar latar holo/aurora pagi hasil render_png.py (<id>.jpg atau <id>.png), kotak kaca tembus pandang, teks asli yang bisa diedit.

Keduanya: catatan pembicara dari "catatan", transisi Fade, mawar pojok kecil di slide isi (--pojok / bawaan mode mandiri),
--animasi: mawar sampul & penutup masuk dengan animasi Zoom bawaan PowerPoint (otomatis sesudah transisi).
Butuh: python-pptx.
"""
import argparse, copy, json, os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ap = argparse.ArgumentParser()
ap.add_argument('--peta', required=True); ap.add_argument('--keluar', required=True)
ap.add_argument('--referensi'); ap.add_argument('--mawar'); ap.add_argument('--latar-folder'); ap.add_argument('--aset')
ap.add_argument('--pojok', action='store_true', help='mode referensi: tempel mawar kecil di pojok slide isi')
ap.add_argument('--animasi', action='store_true', help='mawar sampul & penutup masuk dengan animasi Zoom bawaan PowerPoint')
a = ap.parse_args()
peta = json.load(open(a.peta, encoding='utf-8'))
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

# ---------------------------------------------------------------- umum
def catatan(slide, teks):
    if teks: slide.notes_slide.notes_text_frame.text = teks

def transisi_fade(slide):
    sld = slide._element
    for t in sld.findall(qn('p:transition')): sld.remove(t)
    t = etree.SubElement(sld, qn('p:transition')); t.set('spd', 'med'); etree.SubElement(t, qn('p:fade'))
    # p:transition harus sebelum p:timing / p:extLst
    for tag in ('p:timing', 'p:extLst'):
        el = sld.find(qn(tag))
        if el is not None: el.addprevious(t); break

def tempel_mawar(slide, x, y, w, nama='Mawar'):
    if a.mawar and os.path.exists(a.mawar):
        g = slide.shapes.add_picture(a.mawar, Inches(x), Inches(y), Inches(w), Inches(w)); g.name = nama; return g

def animasi_zoom(slide, shape, dur=1400):
    """Animasi masuk bawaan PowerPoint "Zoom" (preset 53, dari tengah objek) — mulai otomatis sesudah transisi.
    Hanya dipasang bila slide belum punya animasi (animasi asli referensi tidak ditimpa)."""
    sld = slide._element
    if shape is None or sld.find(qn('p:timing')) is not None: return False
    sid = shape.shape_id
    tgt = f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>'
    ukur = lambda i, at: (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{i}" dur="{dur}" fill="hold"/>{tgt}'
                         f'<p:attrNameLst><p:attrName>{at}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav>'
                         f'<p:tav tm="100000"><p:val><p:strVal val="#{at}"/></p:val></p:tav></p:tavLst></p:anim>')
    xml = (f'<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
           f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst>'
           f'<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="5" presetID="53" presetClass="entr" presetSubtype="16" fill="hold" nodeType="afterEffect">'
           f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
           f'<p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
           f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
           + ukur(7, 'ppt_w') + ukur(8, 'ppt_h') +
           f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="9" dur="{dur}"/>{tgt}</p:cBhvr></p:animEffect>'
           f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
           f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml); ext = sld.find(qn('p:extLst'))
    if ext is not None: ext.addprevious(el)
    else: sld.append(el)
    return True

def bersihkan_rel(slide):
    """lepas relasi gambar yang tak lagi dipakai (mis. logo lama yang diganti mawar) agar berkas tidak membengkak"""
    xml = etree.tostring(slide._element).decode()
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith('/image') and f'"{rid}"' not in xml: slide.part.drop_rel(rid)

def cari_bentuk(bentuk, nama):
    for s in bentuk:
        if s.name == nama: return s
        if s.shape_type == 6:  # grup
            r = cari_bentuk(s.shapes, nama)
            if r is not None: return r
    return None

def isi_teks(shape, nilai):
    """Ganti teks, pertahankan format paragraf & run pertama (font, ukuran, warna, poin)."""
    baris = nilai if isinstance(nilai, list) else str(nilai).split('\n')
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    contoh = copy.deepcopy(ps[min(1, len(ps) - 1)] if isinstance(nilai, list) and len(ps) > 1 else ps[0])
    for p in ps: tb.remove(p)
    for t in baris:
        p = copy.deepcopy(contoh); runs = p.findall(qn('a:r'))
        for el in p.findall(qn('a:br')) + p.findall(qn('a:fld')) + runs[1:]: p.remove(el)
        if runs: runs[0].find(qn('a:t')).text = t
        else:  # paragraf kosong: run baru mewarisi format endParaRPr, diletakkan sebelum endParaRPr
            r = etree.Element(qn('a:r')); epr = p.find(qn('a:endParaRPr'))
            if epr is not None: rpr = copy.deepcopy(epr); rpr.tag = qn('a:rPr'); r.append(rpr)
            etree.SubElement(r, qn('a:t')).text = t
            if epr is not None: epr.addprevious(r)
            else: p.append(r)
        tb.append(p)

# ---------------------------------------------------------------- mode referensi
def duplikat(prs, src):
    """Salin slide src utuh ke slide baru di akhir, termasuk gambar & latar (rId dipetakan ulang)."""
    baru = prs.slides.add_slide(src.slide_layout)
    for s in list(baru.shapes): s._element.getparent().remove(s._element)
    peta_rid = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('/notesSlide') or rel.reltype.endswith('/slideLayout'): continue
        peta_rid[rid] = baru.part.relate_to(rel.target_ref if rel.is_external else rel.target_part, rel.reltype, is_external=rel.is_external)
    sumber = src._element.cSld; tujuan = baru._element.cSld
    if sumber.bg is not None:
        if tujuan.bg is not None: tujuan.remove(tujuan.bg)
        tujuan.insert(0, copy.deepcopy(sumber.bg))
    for el in sumber.spTree.iterchildren():
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')): continue
        baru.shapes._spTree.append(copy.deepcopy(el))
    for tag in ('p:transition', 'p:timing'):     # transisi & animasi asli ikut (urutan skema: transition lalu timing)
        el = src._element.find(qn(tag))
        if el is not None: baru._element.append(copy.deepcopy(el))
    for el in baru._element.iter():
        for at in list(el.attrib):
            if at.startswith('{%s}' % RNS) and el.attrib[at] in peta_rid: el.attrib[at] = peta_rid[el.attrib[at]]
    return baru

def mode_referensi():
    prs = Presentation(a.referensi); asli = list(prs.slides); n_asli = len(asli)
    for i, sp in enumerate(peta['slide'], 1):
        pola = sp.get('pola')
        if not pola or not 1 <= pola <= n_asli: sys.exit(f'slide {i}: "pola" harus nomor slide referensi 1..{n_asli}')
        s = duplikat(prs, asli[pola - 1])
        for nama, nilai in (sp.get('isi') or {}).items():
            b = cari_bentuk(s.shapes, nama)
            if b is None or not b.has_text_frame: print(f'PERINGATAN slide {i}: bentuk teks "{nama}" tidak ada di pola {pola}', file=sys.stderr); continue
            isi_teks(b, nilai)
        if sp.get('ganti_gambar'):
            g = cari_bentuk(s.shapes, sp['ganti_gambar'])
            if g is not None:
                w = min(g.width, g.height); mw = tempel_mawar(s, Emu(g.left + (g.width - w) // 2).inches, Emu(g.top + (g.height - w) // 2).inches, Emu(w).inches)
                g._element.getparent().remove(g._element)
                if a.animasi and not animasi_zoom(s, mw): print(f'catatan slide {i}: pola sudah beranimasi, animasi mawar tidak ditambah', file=sys.stderr)
        if isinstance(sp.get('mawar'), dict):
            m = sp['mawar']; mw = tempel_mawar(s, m['x'], m['y'], m['w'])
            if a.animasi: animasi_zoom(s, mw)
        for nama in sp.get('buang', []):
            b = cari_bentuk(s.shapes, nama)
            if b is not None: b._element.getparent().remove(b._element)
        if a.pojok and sp.get('tipe') not in ('sampul', 'penutup'):
            pj = peta.get('pojok', {}); tempel_mawar(s, pj.get('x', Emu(prs.slide_width).inches - .98), pj.get('y', .14), pj.get('w', .8), 'Mawar pojok')
        bersihkan_rel(s)
        if s._element.find(qn('p:transition')) is None: transisi_fade(s)
        catatan(s, sp.get('catatan'))
    lst = prs.slides._sldIdLst
    for sld in list(lst)[:n_asli]: prs.part.drop_rel(sld.rId); lst.remove(sld)
    prs.save(a.keluar)

# ---------------------------------------------------------------- mode mandiri
W, H = 13.333, 7.5
FONT = peta.get('font', 'Segoe UI')
TOKEN = {}
if a.aset:
    for f in ('kc-latar.js', 'kc-latar-hp.js'):
        p = os.path.join(a.aset, f)
        if os.path.exists(p):
            for m in re.finditer(r"'([\w-]+)':\{nama:'[^']*',kel:'\w+',mode:'(\w+)',kaca:'\w+',aksen:'(#[0-9a-fA-F]{6})',aksen2:'(#[0-9a-fA-F]{6})'(?:,teks:'(#[0-9a-fA-F]{6})')?", open(p, encoding='utf-8').read()):
                TOKEN[m.group(1)] = {'gelap': m.group(2) == 'gelap', 'aksen': m.group(3), 'aksen2': m.group(4), 'teks': m.group(5)}
rgb = lambda h: RGBColor.from_string(h.lstrip('#')[:6])

def alpha(fill_el, nilai):
    """beri transparansi pada solidFill (nilai 0–100, persen kepekatan)"""
    clr = fill_el.find('.//' + qn('a:srgbClr'))
    if clr is not None:
        for x in clr.findall(qn('a:alpha')): clr.remove(x)
        etree.SubElement(clr, qn('a:alpha')).set('val', str(int(nilai * 1000)))

def kaca(slide, x, y, w, h, bulat=.18, pekat=58, gelap=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = bulat
    s.fill.solid(); s.fill.fore_color.rgb = rgb('#1c1a33' if gelap else '#ffffff'); alpha(s._element.spPr, pekat if not gelap else 45)
    s.line.color.rgb = rgb('#ffffff'); s.line.width = Pt(1.25); alpha(s._element.spPr.find(qn('a:ln')), 75)
    s.shadow.inherit = False; s.text_frame.text = ''; tanpa_gaya(s)
    return s

def tanpa_gaya(s):
    """buang p:style bawaan (effectRef bayangan tema membuat kaca tampak abu-abu di LibreOffice/PowerPoint)"""
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)

def bulatan(slide, x, y, d, isi, warna, ukuran):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); o.fill.solid(); o.fill.fore_color.rgb = rgb(warna); o.line.fill.background()
    tanpa_gaya(o); tf = o.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = isi; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.runs[0]
    r.font.size = Pt(ukuran); r.font.bold = True; r.font.color.rgb = rgb('#ffffff'); r.font.name = FONT
    return o

def teks(slide, x, y, w, h, isi, ukuran, warna, tebal=False, rata=PP_ALIGN.LEFT, jarak=1.0, spasi_huruf=None, tengah=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.05); tf.vertical_anchor = MSO_ANCHOR.MIDDLE if tengah else MSO_ANCHOR.TOP
    baris = isi if isinstance(isi, list) else [isi]
    for i, t in enumerate(baris):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = rata; p.line_spacing = jarak
        if isinstance(isi, list) and len(baris) > 1: p.space_after = Pt(10)
        r = p.add_run(); r.text = ('•  ' + t) if isinstance(isi, list) and len(baris) > 1 else t
        r.font.size = Pt(ukuran); r.font.bold = tebal; r.font.color.rgb = rgb(warna); r.font.name = FONT
        if spasi_huruf: r.font._element.set('spc', str(spasi_huruf))
    return tb

def mode_mandiri():
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    kosong = prs.slide_layouts[6]
    for i, sp in enumerate(peta['slide'], 1):
        s = prs.slides.add_slide(kosong); tipe = sp.get('tipe', 'poin'); lid = sp.get('latar', peta.get('latar', 'pagi-langit'))
        tk = TOKEN.get(lid, {}); gelap = tk.get('gelap', False)
        T = '#f6f4ff' if gelap else (tk.get('teks') or '#15121e'); T2 = '#d9d4f0' if gelap else '#4a4558'; AK = tk.get('aksen', '#7c3aed')
        gb = next((os.path.join(a.latar_folder, lid + e) for e in ('.jpg', '.png') if a.latar_folder and os.path.exists(os.path.join(a.latar_folder, lid + e))), None)
        if gb: s.shapes.add_picture(gb, 0, 0, prs.slide_width, prs.slide_height)
        else: bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb('#1c1a33' if gelap else '#f4f1fb')
        label = sp.get('label', '')
        if tipe == 'sampul':
            mw = tempel_mawar(s, .25, .55, 6.4)
            if a.animasi: animasi_zoom(s, mw)
            if label: kaca(s, 7.0, 1.9, min(5.8, .14 * len(label) + .7), .5, .5); teks(s, 7.0, 1.9, min(5.8, .14 * len(label) + .7), .5, label.upper(), 11, T, True, PP_ALIGN.CENTER, spasi_huruf=200, tengah=True)
            teks(s, 6.95, 2.55, 6.1, 2.4, sp.get('judul', ''), 60, T, True, jarak=.9)
            teks(s, 7.0, 4.95, 5.9, 1.3, sp.get('sub', ''), 18, T2, jarak=1.2)
        elif tipe == 'penutup':
            mw = tempel_mawar(s, (W - 3.9) / 2, .35, 3.9)
            if a.animasi: animasi_zoom(s, mw)
            teks(s, .8, 4.45, W - 1.6, 1.3, sp.get('judul', 'Terima kasih.'), 56, T, True, PP_ALIGN.CENTER)
            teks(s, .8, 5.75, W - 1.6, .6, sp.get('sub', ''), 18, T2, rata=PP_ALIGN.CENTER)
        elif tipe in ('kelompok', 'pernyataan'):
            teks(s, .8, 2.05, W - 1.6, .4, label.upper(), 11, T2, True, PP_ALIGN.CENTER, spasi_huruf=300)
            teks(s, .8, 2.5, W - 1.6, 1.9, sp.get('judul', ''), 80 if tipe == 'kelompok' else 66, T, True, PP_ALIGN.CENTER, jarak=.95)
            info = sp.get('info') or ([sp['sub']] if sp.get('sub') else [])
            if tipe == 'kelompok':
                lebar = [min(5.6, .12 * len(t) + .9) for t in info]; x = (W - sum(lebar) - .25 * (len(info) - 1)) / 2
                for t, lw in zip(info, lebar): kaca(s, x, 4.8, lw, .62, .5, gelap=gelap); teks(s, x, 4.8, lw, .62, t, 15, T, True, PP_ALIGN.CENTER, tengah=True); x += lw + .25
            else: teks(s, 1.5, 4.6, W - 3, 1.2, info, 20, T2, rata=PP_ALIGN.CENTER, jarak=1.2)
        else:
            teks(s, .85, .78, W - 1.7, .35, label.upper(), 11, T2, True, spasi_huruf=300)
            teks(s, .8, 1.08, W - 1.6, 1.0, sp.get('judul', ''), 40, T, True)
            if tipe == 'anggota':
                ag = sp.get('anggota') or peta.get('anggota', []); per = 3; pw, ph = 3.75, .95
                for j, ang in enumerate(ag):
                    baris, kol = divmod(j, per); n_baris = min(per, len(ag) - baris * per)
                    x = (W - (n_baris * pw + (n_baris - 1) * .3)) / 2 + kol * (pw + .3); y = 2.9 + baris * (ph + .35)
                    kaca(s, x, y, pw, ph, .5, gelap=gelap)
                    bulatan(s, x + .17, y + .15, .65, ''.join(w[0] for w in ang['nama'].split()[:2]).upper(), AK, 15)
                    teks(s, x + .95, y + .12, pw - 1.1, .42, ang['nama'], 16, T, True)
                    teks(s, x + .95, y + .5, pw - 1.1, .35, 'NIM ' + ang['nim'], 12, T2)
            elif tipe == 'agenda':
                it = sp.get('butir') or []; per = 3 if len(it) > 4 else len(it) or 1; pw, ph = 3.6, .9
                for j, t in enumerate(it):
                    baris, kol = divmod(j, per); n_baris = min(per, len(it) - baris * per)
                    x = (W - (n_baris * pw + (n_baris - 1) * .3)) / 2 + kol * (pw + .3); y = 2.75 + baris * (ph + .35)
                    kaca(s, x, y, pw, ph, .5, gelap=gelap)
                    bulatan(s, x + .15, y + .13, .64, str(j + 1), AK, 16)
                    teks(s, x + .95, y, pw - 1.1, ph, t, 20, T, True, tengah=True)
            elif tipe == 'kisi':
                it = sp.get('butir') or []; per = 4; gap = .28; kw = (W - 1.6 - gap * (per - 1)) / per; kh = 1.55
                for j, k in enumerate(it):
                    baris, kol = divmod(j, per); n_baris = min(per, len(it) - baris * per)
                    x = (W - (n_baris * kw + (n_baris - 1) * gap)) / 2 + kol * (kw + gap); y = 2.6 + baris * (kh + .3)
                    kaca(s, x, y, kw, kh, .14, gelap=gelap)
                    teks(s, x + .22, y + .14, kw - .4, .45, k.get('judul', ''), 18, T, True)
                    teks(s, x + .22, y + .6, kw - .4, .5, k.get('rumus', ''), 20, AK, True)
                    if k.get('teks'): teks(s, x + .22, y + 1.08, kw - .4, .4, k['teks'], 12, T2)
            elif tipe in ('kartu', 'banding'):
                kt = sp.get('kartu') or sp.get('kolom') or []; n = max(1, len(kt)); gap = .35; kw = (W - 1.6 - gap * (n - 1)) / n
                for j, k in enumerate(kt):
                    kh = 2.3 if max((len(k.get('teks', '')) for k in kt), default=0) < 110 else 3.0
                    x = .8 + j * (kw + gap); kaca(s, x, 2.75, kw, kh, .12, gelap=gelap)
                    teks(s, x + .3, 2.98, kw - .6, .6, k.get('judul', ''), 24, T, True)
                    teks(s, x + .3, 3.65, kw - .6, 2.0, k.get('teks', ''), 16, T2, jarak=1.2)
            else:
                isi = sp.get('poin') or sp.get('sub', ''); n = len(isi) if isinstance(isi, list) else 2
                kaca(s, .8, 2.55, W - 1.6, min(3.9, .75 + .62 * n), .08, gelap=gelap)
                teks(s, 1.15, 2.85, W - 2.3, min(3.4, .3 + .62 * n), isi, 20, T, jarak=1.15)
        if tipe not in ('sampul', 'penutup'): tempel_mawar(s, W - .98, .14, .8, 'Mawar pojok')   # mawar pojok kanan atas
        transisi_fade(s); catatan(s, sp.get('catatan'))
    prs.save(a.keluar)

if a.referensi: mode_referensi()
else: mode_mandiri()
print(a.keluar, f'{os.path.getsize(a.keluar) // 1024} KB', '(mode referensi)' if a.referensi else '(mode mandiri)')
